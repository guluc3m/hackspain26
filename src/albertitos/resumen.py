"""Resumen Ejecutivo para Alberto — «¿qué pago hoy y por qué?» (T26, BONUS).

Genera `resumen_alberto.html` (y `.pdf` vía typst si el binario está
disponible) desde el store REAL del lote (SOLO LECTURA) + el maestro
(Pedidos_2026 con Importe_Total; NUNCA Pedidos_2025_OLD/backup/NO_TOCAR).

Reglas duras:
- Datos medidos, cero hardcodeo: los totales son la suma exacta de los
  importes del maestro para los PAGAR; lo que no esté en el maestro se
  declara (avisos), jamás se inventa. Sin datos ⇒ PENDIENTE, no ceros falsos.
- Todo con etiqueta medido/estimado y fuentes al pie.
- Lenguaje llano: los códigos de regla se traducen a frases de Alberto.

    uv run python -m albertitos.resumen [--store-root .sdd/lote1] \
        [--maestro caja-de-alberto/FINAL_v7_DEFINITIVO_ahorasi.xlsx] \
        [--salida resumen_alberto]
"""

from __future__ import annotations

import shutil
import sqlite3
import subprocess
import time
from pathlib import Path
from typing import Any

RESULTADOS = ("PAGAR", "NO_PAGAR", "ESCALAR")

# Códigos de regla → frase en llano para Alberto (AGENTS.md §9: sin jerga).
MOTIVOS_LLANO = {
    "ORDER_AMOUNT_MATCHES": "el importe no coincide con el pedido",
    "NO_DOUBLE_PAYMENT": "ya consta un pago para ese pedido (duplicado)",
    "NIF_IN_MASTER": "el proveedor no está dado de alta en el maestro",
    "PROVEEDOR_FANTASMA": "proveedor fantasma: NIF desconocido e IBAN compartido",
    "TOTALS_MUST_MATCH": "los totales de la factura no cuadran",
    "IVA_CONSISTENT": "el IVA declarado no cuadra con el desglose",
    "IBAN_MATCHES_MASTER": "el IBAN no coincide con el del maestro",
    "ORDER_BELONGS_TO_SUPPLIER": "el pedido pertenece a otro proveedor",
    "DATE_VALID_NOT_FUTURE": "la fecha de la factura no se pudo verificar",
    "NO_EMBEDDED_INSTRUCTIONS": "la factura trae instrucciones incrustadas (nunca se siguen)",
    "PEDIDO_EN_REVISION": "el pedido ya está en la cola de revisión",
    "RUNNER_TIMEOUT": "el archivo tardó demasiado y quedó sin leer del todo",
    "ORDER_PENDING": "el pedido no consta como pendiente en el ERP",
}


def _decisions_ro(store_root: Path) -> list[dict[str, Any]]:
    """Lee las decisiones del store REAL en SOLO LECTURA (sqlite modo ro)."""
    db = Path(store_root) / "store.db"
    if not db.is_file():
        return []
    conn = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    try:
        return [dict(r) for r in conn.execute("SELECT * FROM invoices")]
    finally:
        conn.close()


def _maestro(maestro_path: Path) -> tuple[dict[str, float], dict[str, str], list[str]]:
    """(pedido→importe, proveedorID→razón social, avisos de hoja trampa)."""
    from openpyxl import load_workbook

    wb = load_workbook(maestro_path, read_only=True, data_only=True)
    importes: dict[str, float] = {}
    proveedores: dict[str, str] = {}
    avisos: list[str] = []
    hojas_trampa = ("NO_TOCAR", "backup_marzo", "Pedidos_2025_OLD", "v6_deprecated", "MACROS_ROTAS")
    if "Pedidos_2026" not in wb.sheetnames:
        raise ValueError(f"el maestro {maestro_path} no tiene la hoja Pedidos_2026")
    ws = wb["Pedidos_2026"]
    for fila in ws.iter_rows(min_row=2, values_only=True):
        pedido, importe = fila[0], fila[3]
        if pedido and isinstance(importe, (int, float)):
            importes[str(pedido).strip()] = float(importe)
    if "Proveedores" in wb.sheetnames:
        for fila in wb["Proveedores"].iter_rows(min_row=2, values_only=True):
            if fila[0] and fila[1]:
                proveedores[str(fila[0]).strip()] = str(fila[1]).strip()
    for hoja in wb.sheetnames:
        if hoja in hojas_trampa:
            avisos.append(f"hoja trampa ignorada del maestro: {hoja}")
    wb.close()
    return importes, proveedores, avisos


def _codigos(decision: dict[str, Any]) -> list[tuple[str, str]]:
    """rule_codes "CODE:OUTCOME,..." → [(code, outcome), ...]."""
    out = []
    for par in str(decision.get("rule_codes", "")).split(","):
        if ":" in par:
            code, outcome = par.split(":", 1)
            out.append((code.strip(), outcome.strip().upper()))
    return out


def _fallas(decision: dict[str, Any]) -> list[str]:
    return [c for c, o in _codigos(decision) if o == "FAIL"]


def _desconocidos(decision: dict[str, Any]) -> list[str]:
    return [c for c, o in _codigos(decision) if o == "UNKNOWN"]


def datos_resumen(store_root: Path | str, maestro_path: Path | str) -> dict[str, Any]:
    """Agrega el resumen ejecutivo desde el store + maestro (todo medido)."""
    decisiones = _decisions_ro(store_root)
    importes, _proveedores, avisos_maestro = _maestro(maestro_path)

    pagados: list[dict[str, Any]] = []
    sin_importe: list[str] = []
    n_pagar = 0
    for d in decisiones:
        if d["result"] != "PAGAR":
            continue
        n_pagar += 1
        pedido = str(d.get("pedido") or "").strip()
        importe = importes.get(pedido) if pedido else None
        if importe is None:
            sin_importe.append(d["file_id"])
            continue
        pagados.append(
            {"file_id": d["file_id"], "pedido": pedido, "importe_eur": importe}
        )
    pagados.sort(key=lambda x: -x["importe_eur"])
    total_eur = round(sum(x["importe_eur"] for x in pagados), 2)

    no_pago: dict[str, dict[str, Any]] = {}
    n_no_pago = 0
    for d in decisiones:
        if d["result"] != "NO_PAGAR":
            continue
        n_no_pago += 1
        for code in _fallas(d):
            motivo = MOTIVOS_LLANO.get(code, code)
            grupo = no_pago.setdefault(motivo, {"codigo": code, "n": 0, "ejemplos": []})
            grupo["n"] += 1
            if len(grupo["ejemplos"]) < 3:
                grupo["ejemplos"].append(d["file_id"])

    por_revisar: list[dict[str, Any]] = []
    for d in decisiones:
        if d["result"] != "ESCALAR":
            continue
        pedido = str(d.get("pedido") or "").strip()
        importe = importes.get(pedido) if pedido else None
        desconocidos = _desconocidos(d)
        motivo = MOTIVOS_LLANO.get(desconocidos[0], desconocidos[0]) if desconocidos else "revisión requerida"
        por_revisar.append(
            {
                "file_id": d["file_id"],
                "importe_eur": importe,  # None ⇒ «sin importe en el maestro»
                "motivo": motivo,
                "motivo_llano": True,
            }
        )
    por_revisar.sort(key=lambda x: -(x["importe_eur"] or -1))

    # ---- BONUS (T30): plan de revisión por DINERO EN RIESGO.
    # Escalados ordenados por importe desc con suma acumulada y % cubierto;
    # aditivo, sin tocar el motor (la prioridad es presentación para Alberto).
    con_importe = [x for x in por_revisar if x["importe_eur"] is not None]
    sin_importe_n = len(por_revisar) - len(con_importe)
    riesgo_total = round(sum(x["importe_eur"] for x in con_importe), 2)
    plan: list[dict[str, Any]] = []
    acumulado = 0.0
    n_para_80 = None
    umbral = 0.80 * riesgo_total if riesgo_total else None
    for x in con_importe:
        acumulado = round(acumulado + x["importe_eur"], 2)
        plan.append(
            {
                "file_id": x["file_id"],
                "importe_eur": x["importe_eur"],
                "acumulado_eur": acumulado,
                "pct_acumulado": round(100 * acumulado / riesgo_total, 1) if riesgo_total else 0.0,
            }
        )
        if n_para_80 is None and umbral is not None and acumulado >= umbral:
            n_para_80 = len(plan)
    n_para_80 = n_para_80 if n_para_80 is not None else len(plan)

    duplicados = [d["file_id"] for d in decisiones if "NO_DOUBLE_PAYMENT:FAIL" in str(d.get("rule_codes", ""))]
    fantasmas = [
        d["file_id"]
        for d in decisiones
        if "PROVEEDOR_FANTASMA" in str(d.get("rule_codes", ""))
        and ("NIF_IN_MASTER:FAIL" in str(d.get("rule_codes", "")))
    ]
    sin_texto = [
        d["file_id"]
        for d in decisiones
        if d["result"] == "ESCALAR" and "RUNNER_TIMEOUT" in str(d.get("rule_codes", ""))
    ]

    def pareja(valor: Any, etiqueta: str) -> dict[str, Any]:
        return {"valor": str(valor), "etiqueta": etiqueta}

    return {
        "generado": time.strftime("%Y-%m-%d %H:%M"),
        "fuentes": {
            "store": str(Path(store_root) / "store.db"),
            "maestro": str(maestro_path),
            "nota": "SOLO LECTURA — el resumen no decide ni toca el store",
        },
        "pago": {
            "n": {"valor": str(n_pagar), "etiqueta": "medido"},
            "total_eur": (
                {"valor": f"{total_eur:,.2f} EUR", "etiqueta": "medido"}
                if pagados
                else {"valor": "PENDIENTE (sin importes en el maestro)", "etiqueta": "sin datos"}
            ),
            "top10": pagados[:10],
            "sin_importe": sin_importe,
        },
        "no_pago": {
            "n": {"valor": str(n_no_pago), "etiqueta": "medido"},
            "motivos": dict(sorted(no_pago.items(), key=lambda kv: -kv[1]["n"])),
        },
        "revisar": {
            "n": {"valor": str(len(por_revisar)), "etiqueta": "medido"},
            "lista": por_revisar,
        },
        "riesgo": {
            "total_en_riesgo_eur": (
                {"valor": f"{riesgo_total:,.2f} EUR", "etiqueta": "medido"}
                if con_importe
                else {"valor": "PENDIENTE (sin importes en el maestro)", "etiqueta": "sin datos"}
            ),
            "sin_importe": sin_importe_n,
            "plan": plan,
            "para_cubrir_80_pct": {
                "valor": str(n_para_80) if riesgo_total else "PENDIENTE",
                "etiqueta": "medido" if riesgo_total else "sin datos",
                "nota": "mínimo de facturas a revisar (por importe) para cubrir el 80 % del dinero en riesgo",
            },
        },
        "avisos": {
            "duplicados": duplicados,
            "proveedores_fantasma": fantasmas,
            "sin_texto": sin_texto,
            "maestro": avisos_maestro,
    },
    }


def _hoja_html(datos: dict[str, Any]) -> str:
    """Renderiza el dict a HTML (jinja2, mismo estilo llano que la UI)."""
    from jinja2 import Environment, FileSystemLoader

    tpl = FileSystemLoader(str(Path(__file__).parent / "resumen" / "templates"))
    env = Environment(loader=tpl, autoescape=True)
    return env.get_template("resumen.html").render(**datos)


def _typ(datos: dict[str, Any]) -> str:
    """Renderiza el dict a Typst (compilable con el mismo binario del informe)."""
    lineas = [
        "#set page(margin: 2cm)",
        "#set text(size: 10.5pt)",
        "= Resumen para Alberto — ¿qué pago hoy y por qué?",
        f"Generado: {datos['generado']} · fuentes: {datos['fuentes']['store']} (SOLO LECTURA)",
        "",
        "== 1 · Hoy se pagan",
        f"Facturas: {datos['pago']['n']['valor']} (#text(fill: green)[{datos['pago']['n']['etiqueta']}])",
        f"TOTAL: {datos['pago']['total_eur']['valor']} ({datos['pago']['total_eur']['etiqueta']})",
        "",
        "Top 10 por importe:",
        "",
        "#table(",
        "  columns: (auto, 2fr, 1fr),",
        "  table.header([Nº], [Factura · pedido], [Importe]),",
    ]
    for i, x in enumerate(datos["pago"]["top10"], start=1):
        lineas.append(
            f"  [{i}], [{x['file_id']} · pedido {x['pedido']}],"
            f" [{x['importe_eur']:.2f} EUR],"
        )
    lineas.append(")")
    if datos["pago"]["sin_importe"]:
        lineas.append(
            f"NOTA: {len(datos['pago']['sin_importe'])} factura(s) PAGAR sin importe "
            "en el maestro (no se suman; pendiente de revisar el pedido)."
        )
    lineas += [
        "",
        "== 2 · No se pagan",
        f"Total NO_PAGAR: {datos['no_pago']['n']['valor']} ({datos['no_pago']['n']['etiqueta']})",
    ]
    for motivo, g in datos["no_pago"]["motivos"].items():
        lineas.append(f"- {motivo} — {g['n']} factura(s). Ejemplos: {', '.join(g['ejemplos'])}")
    lineas += ["", "== 3 · Te toca mirar (ESCALAR)"]
    lineas.append(f"Total: {datos['revisar']['n']['valor']} (ordenadas por importe). En la cola de Revisión de la UI.")
    for x in datos["revisar"]["lista"]:
        importe = f"{x['importe_eur']:.2f} EUR" if x["importe_eur"] is not None else "sin importe en el maestro"
        lineas.append(f"- {x['file_id']} — {importe}")
    # BONUS T30: plan por dinero en riesgo
    lineas += [
        "",
        "Plan por dinero en riesgo",
        f"Total en juego: {datos['riesgo']['total_en_riesgo_eur']['valor']} ({datos['riesgo']['total_en_riesgo_eur']['etiqueta']})",
        f"Las {datos['riesgo']['para_cubrir_80_pct']['valor']} primeras por importe cubren el 80 % del riesgo.",
        "",
        "#table(",
        "  columns: (auto, 2fr, 1fr, 1fr, 1fr),",
        "  table.header([Nº], [Factura], [Importe], [Acumulado], [% riesgo]),",
    ]
    for i, x in enumerate(datos["riesgo"]["plan"][:10], start=1):
        lineas.append(
            f"  [{i}], [{x['file_id']}], [{x['importe_eur']:.2f} EUR],"
            f" [{x['acumulado_eur']:.2f} EUR], [{x['pct_acumulado']} %],"
        )
    lineas.append(")")
    lineas += ["", "== 4 · Avisos"]
    lineas.append(f"- Duplicados detectados (NO_DOUBLE_PAYMENT): {len(datos['avisos']['duplicados'])}")
    lineas.append(f"- Proveedores fantasma (no dados de alta, IBAN compartido): {len(datos['avisos']['proveedores_fantasma'])}")
    lineas.append(f"- Sin texto legible / demasiado lentas (escaladas): {len(datos['avisos']['sin_texto'])}")
    for a in datos["avisos"]["maestro"]:
        lineas.append(f"- {a}")
    return "\n".join(lineas) + "\n"


def _binario_typst() -> Path | None:
    cual = shutil.which("typst")
    if cual:
        return Path(cual)
    candidato = Path.home() / ".local" / "bin" / "typst"
    return candidato if candidato.is_file() else None


def generar_resumen(
    store_root: Path | str,
    maestro_path: Path | str,
    salida: Path | str | None = None,
) -> dict[str, Path]:
    """Genera resumen_alberto.html y resumen_alberto.pdf (si hay typst)."""
    datos = datos_resumen(store_root, maestro_path)
    base = Path(salida) if salida is not None else Path("resumen_alberto")
    rutas: dict[str, Path] = {}
    html = _hoja_html(datos)
    rutas["html"] = base.with_suffix(".html")
    rutas["html"].write_text(html, encoding="utf-8")
    binario = _binario_typst()
    if binario is not None:
        typ = base.with_suffix(".typ")
        typ.write_text(_typ(datos), encoding="utf-8")
        pdf = base.with_suffix(".pdf")
        comp = subprocess.run(
            [str(binario), "compile", str(typ), str(pdf)],
            capture_output=True, text=True, check=False,
        )
        if comp.returncode == 0:
            rutas["pdf"] = pdf
        rutas["typ"] = typ
    return rutas


def main(argv: list[str] | None = None) -> int:
    import argparse

    parser = argparse.ArgumentParser(
        prog="albertitos.resumen",
        description="Resumen ejecutivo para Alberto: ¿qué pago hoy y por qué?",
    )
    parser.add_argument("--store-root", default=".sdd/lote1")
    parser.add_argument("--maestro", default="/home/deploy/hackspain26/caja-de-alberto/FINAL_v7_DEFINITIVO_ahorasi.xlsx")
    parser.add_argument("--salida", default="resumen_alberto")
    args = parser.parse_args(argv)
    rutas = generar_resumen(args.store_root, args.maestro, args.salida)
    for formato, ruta in sorted(rutas.items()):
        print(f"{formato}: {ruta.resolve()}")
    if "pdf" not in rutas:
        print("pdf: PENDIENTE — binario typst no disponible (el HTML está completo)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
