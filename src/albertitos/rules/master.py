"""Carga del maestro (caja-de-alberto Excel) — SOLO lectura.

Lee únicamente `Proveedores`, `Pedidos_2026` y `pendiente_revisar`. Las hojas
trampa del yaml (`hojas_ignoradas`) se ignoran explícitamente y quedan
registradas en el snapshot de configuración. La fila P007 duplicada se
deduplica (primera ocurrencia gana) y se deja registrado.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from albertitos.parse.normalizers import normalize_iban, normalize_nif, parse_amount

# Hojas que SÍ se leen (título → lector). Cualquier otra = ignorada.
HOJAS_MAESTRO = ("Proveedores", "Pedidos_2026")
HOJA_REVISION = "pendiente_revisar"

_RE_PEDIDO = re.compile(r"\bPO-\d{4}-\d{3,5}\b")


@dataclass(frozen=True)
class Proveedor:
    id: str
    razon_social: str
    nif: str
    iban: str  # normalizado sin espacios


@dataclass(frozen=True)
class Pedido:
    id: str
    proveedor_id: str
    nif: str
    importe: float
    estado: str
    fecha: str


@dataclass
class Maestro:
    proveedores_por_nif: dict[str, Proveedor]
    proveedores_por_id: dict[str, Proveedor]
    pedidos: dict[str, Pedido]
    pedidos_en_revision: frozenset[str]
    hojas_ignoradas: tuple[str, ...]
    duplicados_deducidos: tuple[str, ...]
    sha256: str
    avisos: tuple[str, ...] = field(default_factory=tuple)


def load_master(
    path: str | Path,
    hojas_ignoradas: list[str] | tuple[str, ...] = (),
) -> Maestro:
    """Carga el Excel del maestro. Nunca escribe en él (submódulo read-only)."""
    path = Path(path)
    raw = path.read_bytes()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    hojas_presentes = set(wb.sheetnames)
    usadas = set(HOJAS_MAESTRO) | {HOJA_REVISION}
    ignoradas = tuple(
        h for h in hojas_ignoradas if h in hojas_presentes
    )
    no_configuradas = tuple(sorted(hojas_presentes - usadas - set(ignoradas)))

    avisos: list[str] = [
        f"hoja_no_leida:{h}" for h in no_configuradas
    ]

    # --- Proveedores: dedup por ID, primera ocurrencia gana
    proveedores_por_nif: dict[str, Proveedor] = {}
    proveedores_por_id: dict[str, Proveedor] = {}
    duplicados: list[str] = []
    if "Proveedores" in hojas_presentes:
        ws = wb["Proveedores"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            pid = str(row[0]).strip()
            prov = Proveedor(
                id=pid,
                razon_social=str(row[1]).strip(),
                nif=normalize_nif(str(row[2] or "")),
                iban=normalize_iban(str(row[3] or "")),
            )
            if pid in proveedores_por_id:
                duplicados.append(pid)
                continue
            proveedores_por_id[pid] = prov
            proveedores_por_nif[prov.nif] = prov
    if duplicados:
        avisos.append(
            "proveedores_deduplicados:" + ",".join(sorted(set(duplicados)))
        )

    # --- Pedidos
    pedidos: dict[str, Pedido] = {}
    if "Pedidos_2026" in hojas_presentes:
        ws = wb["Pedidos_2026"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            pid = str(row[0]).strip()
            # T38-F7: pedido duplicado ⇒ PRIMERA gana + aviso (como Proveedores;
            # el lote 2 puede traer filas repetidas del ERP).
            if pid in pedidos:
                avisos.append("pedido_deduplicado:" + pid)
                continue
            importe = row[3] if row[3] is not None else 0.0
            # T38-F7: el ERP puede traer el importe como TEXTO ("1.234,56")
            # — fallback a parse_amount en vez de caer el runner completo.
            if isinstance(importe, (int, float)):
                importe_num = float(importe)
            else:
                importe_num = parse_amount(str(importe).strip())
                if importe_num is None:
                    avisos.append(f"importe_ilegible:{pid}")
                    importe_num = 0.0
            pedidos[pid] = Pedido(
                id=pid,
                proveedor_id=str(row[1] or "").strip(),
                nif=normalize_nif(str(row[2] or "")),
                importe=importe_num,
                estado=str(row[4] or "").strip().upper(),
                fecha=str(row[5] or "").strip(),
            )

    # --- pendiente_revisar: pedidos marcados para ojo humano ⇒ ESCALAR
    revision: set[str] = set()
    if HOJA_REVISION in hojas_presentes:
        ws = wb[HOJA_REVISION]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell and _RE_PEDIDO.search(str(cell)):
                    revision.add(_RE_PEDIDO.search(str(cell)).group(0))

    wb.close()
    return Maestro(
        proveedores_por_nif=proveedores_por_nif,
        proveedores_por_id=proveedores_por_id,
        pedidos=pedidos,
        pedidos_en_revision=frozenset(revision),
        hojas_ignoradas=ignoradas,
        duplicados_deducidos=tuple(duplicados),
        sha256=hashlib.sha256(raw).hexdigest(),
        avisos=tuple(avisos),
    )